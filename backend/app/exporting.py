"""Turn companies / people into table rows and download them as CSV or Excel (.xlsx).

company_table / people_table return (headers, rows); the same rows feed the CSV and Excel
downloads and the Google Sheets export (app/sheets.py).
"""
import csv
import io
import re
from collections.abc import Iterable

from fastapi.responses import Response, StreamingResponse

from .models import Company, Contact
from .people import SENIORITY_RANK

Cell = str | int | float | None

# Same columns and order as the original CSV export (Steps 1-3), then the Step 4-6 extras.
COMPANY_HEADERS = [
    "name", "category", "phone", "email", "emails", "website", "domain", "address", "city", "country",
    "lat", "lon", "rating", "reviews_count", "opening_hours", "google_maps_url", "business_status",
    "facebook", "instagram", "linkedin", "platform", "https", "copyright_year", "contacts", "source",
    # Step 4-6
    "email_status", "lead_status", "score", "score_reasons",
    "decision_maker", "decision_maker_title", "decision_maker_email",
    # Phone check (app/phones.py)
    "phone_type",
    # Company size (app/sizing.py) and the US PPP loan (app/ppp.py)
    "size_band", "employees", "employees_source", "founded_year", "locations", "hiring",
    "ppp_jobs", "ppp_loan_date",
]

PEOPLE_HEADERS = [
    "name", "first_name", "last_name", "title", "seniority", "decision_maker", "email", "email_status",
    "email_confidence", "phone", "linkedin", "company", "company_website", "company_phone", "city",
    "country", "category", "lead_status", "score",
]


def _yes_no(value) -> str:
    return "" if value is None else ("yes" if value else "no")


def _contacts_text(company: Company) -> str:
    """'Ali Khan (Director) <ali@x.com>; Sara <sara@x.com>'"""
    parts = []
    for c in company.contacts or []:
        text = c.full_name or ""
        if c.title:
            text += f" ({c.title})"
        if c.email:
            text += f" <{c.email}>"
        if text.strip():
            parts.append(text.strip())
    return "; ".join(parts)


def _unique_emails(company: Company) -> list[str]:
    """Best email first, no duplicates (case-insensitive)."""
    seen, result = set(), []
    for e in [company.email, *(company.emails or [])]:
        if e and e.lower() not in seen:
            seen.add(e.lower())
            result.append(e)
    return result


def best_decision_maker(company: Company) -> Contact | None:
    """The most senior decision-maker, preferring people with an email."""
    deciders = [c for c in company.contacts or [] if c.is_decision_maker]
    if not deciders:
        return None
    return min(deciders, key=lambda c: (not c.email, SENIORITY_RANK.get(c.seniority, 9), c.id or 0))


def company_table(companies: Iterable[Company]) -> tuple[list[str], list[list[Cell]]]:
    rows = []
    for c in companies:
        socials = c.socials or {}
        signals = c.website_signals or {}
        crawled = c.enrich_status == "done"  # website facts are only reliable after a good crawl
        # Size facts (hiring) stay true when a later re-crawl fails: same rule as app/sizing.py.
        size_facts = bool(signals.get("pages_crawled"))
        statuses = c.email_statuses or {}
        dm = best_decision_maker(c)
        ppp = c.ppp or {}
        rows.append([
            c.name, c.category, c.phone, c.email, "; ".join(_unique_emails(c)), c.website, c.domain,
            c.address, c.city, c.country, c.lat, c.lon, c.rating, c.reviews_count, c.opening_hours,
            c.google_maps_url, c.business_status,
            socials.get("facebook"), socials.get("instagram"), socials.get("linkedin"),
            signals.get("platform") if crawled else None,
            _yes_no(signals.get("https")) if crawled else "",
            signals.get("copyright_year") if crawled else None,
            _contacts_text(c), c.source,
            statuses.get(c.email) if c.email else None,
            c.lead_status or "new", c.score, "; ".join(c.score_reasons or []),
            dm.full_name if dm else None, dm.title if dm else None, dm.email if dm else None,
            c.phone_type,
            c.size_band, c.employees, c.employees_source, c.founded_year, c.locations_count,
            _yes_no(bool(signals.get("has_careers"))) if size_facts else "",
            ppp.get("jobs"), ppp.get("date_approved"),
        ])
    return list(COMPANY_HEADERS), rows


def people_table(contacts: Iterable[Contact]) -> tuple[list[str], list[list[Cell]]]:
    rows = []
    for p in contacts:
        co = p.company
        rows.append([
            p.full_name, p.first_name, p.last_name, p.title, p.seniority, _yes_no(bool(p.is_decision_maker)),
            p.email, p.email_status, p.email_confidence, p.phone, p.linkedin_url,
            co.name if co else None, co.website if co else None, co.phone if co else None,
            co.city if co else None, co.country if co else None, co.category if co else None,
            p.lead_status or "new", p.score,
        ])
    return list(PEOPLE_HEADERS), rows


# --- spreadsheet safety --------------------------------------------------------------

# Text scraped from websites could start with "=" and run as a formula when the file is
# opened in Excel ("CSV injection"). Phone numbers like "+1 555-0100" are left alone.
_SAFE_NUMBERISH = re.compile(r"[+\-]?[\d\s().+\-/]*")


def safe_cell(value: Cell) -> Cell:
    if not isinstance(value, str) or not value:
        return value
    first = value[0]
    if first in "=@\t\r" or (first in "+-" and not _SAFE_NUMBERISH.fullmatch(value)):
        return "'" + value
    return value


def _download_headers(filename: str) -> dict[str, str]:
    return {"Content-Disposition": f'attachment; filename="{filename}"'}


def csv_response(headers: list[str], rows: list[list[Cell]], filename: str) -> StreamingResponse:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if v is None else safe_cell(v) for v in row])
    # BOM so Excel opens Urdu/Hindi characters correctly
    data = "﻿" + buf.getvalue()
    return StreamingResponse(
        iter([data]), media_type="text/csv; charset=utf-8", headers=_download_headers(filename)
    )


LINK_COLUMNS = {"website", "company_website", "google_maps_url", "linkedin", "facebook", "instagram"}
MAX_COLUMN_WIDTH = 50


def xlsx_bytes(headers: list[str], rows: list[list[Cell]], sheet_title: str = "Leads") -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"  # keep the header row visible while scrolling

    widths = [len(h) for h in headers]
    link_font = Font(color="0563C1", underline="single")
    for r, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            if isinstance(value, str):
                # Excel files can't hold control characters; strip them.
                value = safe_cell(ILLEGAL_CHARACTERS_RE.sub("", value))
            cell = ws.cell(row=r, column=col, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # always text, never a formula
                widths[col - 1] = max(widths[col - 1], len(value))
                if headers[col - 1] in LINK_COLUMNS and value.startswith(("http://", "https://")):
                    cell.hyperlink = value
                    cell.font = link_font
            elif value is not None:
                widths[col - 1] = max(widths[col - 1], len(str(value)))
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, MAX_COLUMN_WIDTH)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(headers: list[str], rows: list[list[Cell]], filename: str) -> Response:
    return Response(
        content=xlsx_bytes(headers, rows),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_download_headers(filename),
    )
