"""Table building + CSV/Excel downloads (app/exporting.py)."""
import asyncio
import csv
import io

from openpyxl import load_workbook

from app import exporting
from app.models import Company, Contact


def make_company(**kw) -> Company:
    fields = dict(name="Acme Dental", source="osm", source_ref="x", website="https://acme.example",
                  lead_status="new", contacts=[])
    fields.update(kw)
    return Company(**fields)


def row_dict(headers, row):
    return dict(zip(headers, row))


def test_company_table_keeps_old_csv_columns_first():
    headers, _ = exporting.company_table([])
    old = ["name", "category", "phone", "email", "emails", "website", "domain", "address", "city", "country",
           "lat", "lon", "rating", "reviews_count", "opening_hours", "google_maps_url", "business_status",
           "facebook", "instagram", "linkedin", "platform", "https", "copyright_year", "contacts", "source"]
    assert headers[: len(old)] == old
    assert headers[len(old):] == ["email_status", "lead_status", "score", "score_reasons",
                                  "decision_maker", "decision_maker_title", "decision_maker_email"]


def test_company_row_values():
    staff = Contact(id=2, full_name="Sara", title="Receptionist", email="sara@acme.example")
    owner = Contact(id=1, full_name="Ali Khan", title="Owner", seniority="owner", is_decision_maker=True,
                    email="ali@acme.example")
    c = make_company(
        email="info@acme.example", emails=["Info@acme.example", "sales@acme.example"],
        email_statuses={"info@acme.example": "valid"}, socials={"facebook": "https://fb.com/acme"},
        enrich_status="done", website_signals={"platform": "WordPress", "https": False, "copyright_year": 2019},
        score=55, score_reasons=["Verified email +25", "Phone +10"], contacts=[staff, owner],
    )
    headers, rows = exporting.company_table([c])
    row = row_dict(headers, rows[0])
    assert row["emails"] == "info@acme.example; sales@acme.example"  # no case-duplicate
    assert row["facebook"] == "https://fb.com/acme"
    assert (row["platform"], row["https"], row["copyright_year"]) == ("WordPress", "no", 2019)
    assert row["contacts"] == "Sara (Receptionist) <sara@acme.example>; Ali Khan (Owner) <ali@acme.example>"
    assert row["email_status"] == "valid" and row["lead_status"] == "new" and row["score"] == 55
    assert row["score_reasons"] == "Verified email +25; Phone +10"
    assert (row["decision_maker"], row["decision_maker_title"], row["decision_maker_email"]) == (
        "Ali Khan", "Owner", "ali@acme.example")


def test_failed_crawl_does_not_claim_no_ssl():
    c = make_company(enrich_status="failed", website_signals={"https": False, "platform": None})
    headers, rows = exporting.company_table([c])
    row = row_dict(headers, rows[0])
    assert row["https"] == "" and row["platform"] is None


def test_people_table():
    co = make_company(phone="021-111", city="Karachi", country="Pakistan", category="dentist")
    p = Contact(full_name="Ali Khan", first_name="Ali", last_name="Khan", title="Owner", seniority="owner",
                is_decision_maker=True, email="ali@acme.example", email_status="valid", email_confidence=90,
                lead_status="contacted", score=80, company=co)
    headers, rows = exporting.people_table([p, Contact(full_name="No Company")])
    assert headers == ["name", "first_name", "last_name", "title", "seniority", "decision_maker", "email",
                       "email_status", "email_confidence", "phone", "linkedin", "company", "company_website",
                       "company_phone", "city", "country", "category", "lead_status", "score"]
    row = row_dict(headers, rows[0])
    assert row["decision_maker"] == "yes" and row["company"] == "Acme Dental"
    assert row["company_website"] == "https://acme.example" and row["company_phone"] == "021-111"
    assert row["city"] == "Karachi" and row["lead_status"] == "contacted" and row["score"] == 80
    other = row_dict(headers, rows[1])
    assert other["company"] is None and other["decision_maker"] == "no" and other["lead_status"] == "new"


def test_safe_cell_blocks_formulas_but_keeps_phones():
    assert exporting.safe_cell("=HYPERLINK(\"http://evil\")") == "'=HYPERLINK(\"http://evil\")"
    assert exporting.safe_cell("@SUM(A1)") == "'@SUM(A1)"
    assert exporting.safe_cell("-2+3+cmd|' /C calc'!A0") == "'-2+3+cmd|' /C calc'!A0"
    assert exporting.safe_cell("+1 (555) 010-0100") == "+1 (555) 010-0100"
    assert exporting.safe_cell("Dentist") == "Dentist"
    assert exporting.safe_cell(42) == 42


def _body(resp) -> str:
    """Read a StreamingResponse body in a test."""
    async def collect():
        return "".join([c if isinstance(c, str) else c.decode() async for c in resp.body_iterator])

    return asyncio.run(collect())


def test_csv_response_has_bom_and_blanks():
    resp = exporting.csv_response(["a", "b"], [["=1+1", None], ["Ünïcode", 3]], "x.csv")
    assert resp.headers["content-disposition"] == 'attachment; filename="x.csv"'
    assert _body(resp) == "\ufeffa,b\r\n'=1+1,\r\nÜnïcode,3\r\n"


def test_people_csv_via_api(client, db):
    co = Company(name="=Evil Co", source="osm", source_ref="e", phone="+92 300 1234567")
    co.contacts.append(Contact(full_name="Ali", title="Owner"))
    db.add(co)
    db.commit()
    r = client.get("/api/people/export.csv")
    assert r.status_code == 200 and r.text.startswith("\ufeffname,first_name")
    row = next(csv.DictReader(io.StringIO(r.text.lstrip("\ufeff"))))
    assert row["company"] == "'=Evil Co" and row["company_phone"] == "+92 300 1234567"


def test_xlsx_bytes_styling_and_safety():
    data = exporting.xlsx_bytes(
        ["name", "website", "score"],
        [["=Evil", "https://acme.example", 5], ["Bad\x07char", None, None]],
    )
    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["name", "website", "score"]
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws["A2"].value == "'=Evil" and ws["A2"].data_type == "s"
    assert ws["B2"].hyperlink.target == "https://acme.example"
    assert ws["C2"].value == 5
    assert ws["A3"].value == "Badchar"
    assert ws.column_dimensions["B"].width <= exporting.MAX_COLUMN_WIDTH


def test_csv_rows_parse_back():
    buf = io.StringIO()
    headers, rows = exporting.company_table([make_company(phone="+1 555")])
    csv.writer(buf).writerows([headers, *rows])
    parsed = list(csv.reader(io.StringIO(buf.getvalue())))
    assert parsed[1][headers.index("phone")] == "+1 555"
