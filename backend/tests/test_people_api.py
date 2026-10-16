"""People search API (GET/PATCH /api/people, exports) and POST /api/status/bulk."""
import io

import pytest
from openpyxl import load_workbook

from app.models import Activity, Company, Contact, JobResult, LeadList, ListItem, SearchJob


@pytest.fixture
def people(db):
    """Two companies from one search, a third from another, and five people."""
    job = SearchJob(kind="search", query="dentist", location="Miami", source="osm", status="done")
    other_job = SearchJob(kind="search", query="lawyer", location="Lahore", source="osm", status="done")
    db.add_all([job, other_job])
    db.flush()
    smile = Company(name="Smile Dental", source="osm", source_ref="1", city="Miami", category="dentist",
                    phone="305-1", website="https://smile.example")
    bright = Company(name="Bright Teeth", source="osm", source_ref="2", city="Miami Beach", category="dentist")
    law = Company(name="Khan Law", source="osm", source_ref="3", city="Lahore", category="lawyer")
    db.add_all([smile, bright, law])
    db.flush()
    db.add_all([JobResult(job_id=job.id, company_id=smile.id), JobResult(job_id=job.id, company_id=bright.id),
                JobResult(job_id=other_job.id, company_id=law.id)])
    rows = {
        "ali": Contact(company_id=smile.id, full_name="Ali Owner", title="Owner / CTO", seniority="owner",
                       is_decision_maker=True, email="ali@smile.example", email_status="valid", score=90),
        "sara": Contact(company_id=smile.id, full_name="Sara Front", title="Front Desk Coordinator",
                        seniority="staff", email="sara@smile.example", email_status="catch_all", score=20),
        "omar": Contact(company_id=bright.id, full_name="Omar Manager", title="Office Manager", seniority="manager",
                        is_decision_maker=True, score=50),
        "dan": Contact(company_id=bright.id, full_name="Dan Director", title="Director of Operations",
                       seniority="director", is_decision_maker=True, email="dan@bright.example",
                       email_status="unknown", score=None),
        "zara": Contact(company_id=law.id, full_name="Zara Attorney", title="Attorney", seniority="professional",
                        email="zara@law.example", email_status="invalid", score=30, lead_status="contacted"),
    }
    db.add_all(rows.values())
    db.commit()
    return {"job": job, "smile": smile, "bright": bright, "law": law, **rows}


def names(resp) -> list[str]:
    assert resp.status_code == 200, resp.text
    return [p["full_name"] for p in resp.json()["items"]]


def test_default_sort_is_most_senior_first(client, people):
    r = client.get("/api/people")
    assert names(r) == ["Ali Owner", "Dan Director", "Omar Manager", "Zara Attorney", "Sara Front"]
    body = r.json()
    assert body["total"] == 5
    ali = body["items"][0]
    assert ali["company"]["name"] == "Smile Dental" and ali["company"]["city"] == "Miami"
    assert ali["list_ids"] == [] and ali["is_decision_maker"] is True and ali["lead_status"] == "new"


def test_sort_by_score_and_name(client, people):
    assert names(client.get("/api/people?sort=score")) == [
        "Ali Owner", "Omar Manager", "Zara Attorney", "Sara Front", "Dan Director"]  # no score last
    assert names(client.get("/api/people?sort=name")) == [
        "Ali Owner", "Dan Director", "Omar Manager", "Sara Front", "Zara Attorney"]
    assert client.get("/api/people?sort=bogus").status_code == 422


def test_paging(client, people):
    r = client.get("/api/people?limit=2&offset=2")
    assert names(r) == ["Omar Manager", "Zara Attorney"] and r.json()["total"] == 5


@pytest.mark.parametrize("query, expected", [
    ("title=owner,manager", {"Ali Owner", "Omar Manager"}),
    ("title=cto", {"Ali Owner"}),  # word start: "cto" does not match "Director"
    ("title=OPERATIONS", {"Dan Director"}),
    ("seniority=owner,director", {"Ali Owner", "Dan Director"}),
    ("decision_makers=true", {"Ali Owner", "Omar Manager", "Dan Director"}),
    ("has_email=true", {"Ali Owner", "Sara Front", "Dan Director", "Zara Attorney"}),
    ("has_email=false", {"Omar Manager"}),
    ("email_status=valid,catch_all", {"Ali Owner", "Sara Front"}),
    ("lead_status=contacted", {"Zara Attorney"}),
    ("lead_status=new", {"Ali Owner", "Sara Front", "Omar Manager", "Dan Director"}),
    ("city=miami", {"Ali Owner", "Sara Front", "Omar Manager", "Dan Director"}),
    ("category=law", {"Zara Attorney"}),
    ("min_score=50", {"Ali Owner", "Omar Manager"}),
    ("q=bright", {"Omar Manager", "Dan Director"}),  # company name
    ("q=attorney", {"Zara Attorney"}),  # title
    ("q=100%", set()),  # "%" is a literal, not a wildcard
])
def test_filters(client, people, query, expected):
    assert set(names(client.get(f"/api/people?{query}"))) == expected


def test_job_company_and_list_filters(client, db, people):
    assert set(names(client.get(f"/api/people?job_id={people['job'].id}"))) == {
        "Ali Owner", "Sara Front", "Omar Manager", "Dan Director"}
    assert set(names(client.get(f"/api/people?company_id={people['bright'].id}"))) == {"Omar Manager", "Dan Director"}
    lead_list = LeadList(name="Hot")
    db.add(lead_list)
    db.flush()
    # A company in the list does not bring its people along - only people added directly.
    db.add_all([ListItem(list_id=lead_list.id, contact_id=people["zara"].id),
                ListItem(list_id=lead_list.id, company_id=people["smile"].id)])
    db.commit()
    r = client.get(f"/api/people?list_id={lead_list.id}")
    assert names(r) == ["Zara Attorney"] and r.json()["items"][0]["list_ids"] == [lead_list.id]


def test_patch_person_status_logs_activity_and_rescoring(client, db, people):
    pid = people["omar"].id
    r = client.patch(f"/api/people/{pid}", json={"lead_status": "interested"})
    assert r.status_code == 200
    person = r.json()
    assert person["lead_status"] == "interested" and person["company"]["name"] == "Bright Teeth"
    assert person["score"] == 35 and "Decision-maker +30" in person["score_reasons"]
    acts = db.query(Activity).all()
    assert len(acts) == 1
    assert acts[0].company_id == people["bright"].id and acts[0].contact_id == pid and acts[0].kind == "status"
    assert acts[0].text == "Omar Manager - Status: New → Interested"
    # Same status again: nothing logged
    client.patch(f"/api/people/{pid}", json={"lead_status": "interested"})
    assert db.query(Activity).count() == 1
    # Lost -> score 0
    assert client.patch(f"/api/people/{pid}", json={"lead_status": "lost"}).json()["score"] == 0


def test_patch_person_errors(client, people):
    assert client.patch("/api/people/9999", json={"lead_status": "won"}).status_code == 404
    assert client.patch(f"/api/people/{people['ali'].id}", json={"lead_status": "maybe"}).status_code == 422


def test_people_exports(client, people):
    r = client.get("/api/people/export.csv?decision_makers=true")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    lines = r.text.lstrip("﻿").strip().splitlines()
    assert lines[0].startswith("name,first_name,last_name,title,seniority,decision_maker,email")
    assert [line.split(",")[0] for line in lines[1:]] == ["Ali Owner", "Dan Director", "Omar Manager"]

    r = client.get("/api/people/export.xlsx?city=lahore")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A2"].value == "Zara Attorney" and ws.max_row == 2


def test_bulk_status(client, db, people):
    body = {"company_ids": [people["smile"].id, people["law"].id, 9999],
            "contact_ids": [people["ali"].id, people["zara"].id], "lead_status": "contacted"}
    r = client.post("/api/status/bulk", json=body)
    # smile + law companies, ali (zara already "contacted")
    assert r.status_code == 200 and r.json() == {"changed": 3}
    db.expire_all()
    assert db.get(Company, people["smile"].id).lead_status == "contacted"
    assert db.get(Company, people["smile"].id).score is not None  # rescored
    texts = sorted(a.text for a in db.query(Activity))
    assert texts == ["Ali Owner - Status: New → Contacted", "Status: New → Contacted", "Status: New → Contacted"]
    assert client.post("/api/status/bulk", json={"lead_status": "nope"}).status_code == 422
