"""Companies API (Steps 4-6): new filters, sorting, contacts order, list ids, lead status
with activity log, notes, Excel export and "find emails" for picked companies."""
import csv
import io

import pytest
from openpyxl import load_workbook

from app.models import Activity, Company, Contact, JobResult, LeadList, ListItem, SearchJob


def add_company(db, name, **kw) -> Company:
    c = Company(name=name, source="osm", source_ref=name, **kw)
    db.add(c)
    db.commit()
    return c


def names(resp) -> list[str]:
    assert resp.status_code == 200, resp.text
    return [c["name"] for c in resp.json()["items"]]


@pytest.fixture
def sample(db):
    alpha = add_company(db, "alpha", score=40, rating=4.1, reviews_count=10, lead_status="contacted")
    beta = add_company(db, "Beta", score=90, rating=None)
    gamma = add_company(db, "Gamma", score=None, rating=4.9, reviews_count=200)
    delta = add_company(db, "Delta", score=70, rating=4.1, reviews_count=50, email="info@delta.pk")
    return alpha, beta, gamma, delta


# ----- filters & sorting -------------------------------------------------------------


def test_sort_by_name_score_and_rating(client, sample):
    assert names(client.get("/api/companies")) == ["alpha", "Beta", "Delta", "Gamma"]  # case-insensitive
    assert names(client.get("/api/companies?sort=score")) == ["Beta", "Delta", "alpha", "Gamma"]  # empty last
    assert names(client.get("/api/companies?sort=rating")) == ["Gamma", "Delta", "alpha", "Beta"]
    assert client.get("/api/companies?sort=nope").status_code == 422


def test_filter_lead_status_min_score_and_list(client, db, sample):
    alpha, beta, gamma, delta = sample
    assert names(client.get("/api/companies?lead_status=contacted")) == ["alpha"]
    assert names(client.get("/api/companies?lead_status=new")) == ["Beta", "Delta", "Gamma"]
    assert names(client.get("/api/companies?min_score=70")) == ["Beta", "Delta"]

    hot, cold = LeadList(name="Hot"), LeadList(name="Cold")
    db.add_all([hot, cold])
    db.flush()
    db.add_all([ListItem(list_id=hot.id, company_id=gamma.id), ListItem(list_id=hot.id, company_id=beta.id),
                ListItem(list_id=cold.id, company_id=beta.id)])
    db.commit()
    assert names(client.get(f"/api/companies?list_id={hot.id}")) == ["Beta", "Gamma"]
    assert names(client.get(f"/api/companies?list_id={hot.id}&min_score=50")) == ["Beta"]

    items = {c["name"]: c for c in client.get("/api/companies").json()["items"]}
    assert items["Beta"]["list_ids"] == [hot.id, cold.id]
    assert items["Gamma"]["list_ids"] == [hot.id]
    assert items["alpha"]["list_ids"] == []


def test_filters_usable_as_plain_function(db, sample):
    from app.routers.companies import filters

    stmt = filters(min_score=50, has_email=True)
    assert [c.name for c in db.scalars(stmt)] == ["Delta"]


def test_job_filter(client, db, sample):
    _, beta, *_ = sample
    job = SearchJob(kind="search", query="q", location="l", source="osm", status="done")
    db.add(job)
    db.flush()
    db.add(JobResult(job_id=job.id, company_id=beta.id))
    db.commit()
    assert names(client.get(f"/api/companies?job_id={job.id}")) == ["Beta"]


def test_contacts_sorted_most_senior_first(client, db):
    c = add_company(db, "Clinic")
    db.add_all([
        Contact(company_id=c.id, full_name="Zara Staff", seniority="staff"),
        Contact(company_id=c.id, full_name="No Title"),
        Contact(company_id=c.id, full_name="Mona Manager", seniority="manager", is_decision_maker=False),
        Contact(company_id=c.id, full_name="Owen Owner", seniority="owner", is_decision_maker=True),
        Contact(company_id=c.id, full_name="Adam Manager", seniority="manager", is_decision_maker=False),
        Contact(company_id=c.id, full_name="Omar Office Manager", seniority="manager", is_decision_maker=True),
    ])
    db.commit()
    expected = ["Owen Owner", "Omar Office Manager", "Adam Manager", "Mona Manager", "Zara Staff", "No Title"]
    listed = client.get("/api/companies").json()["items"][0]["contacts"]
    assert [p["full_name"] for p in listed] == expected
    one = client.get(f"/api/companies/{c.id}").json()["contacts"]
    assert [p["full_name"] for p in one] == expected


# ----- lead status, activities & notes ------------------------------------------------


def test_patch_status_logs_activity_and_rescores(client, db):
    c = add_company(db, "Acme", phone="+92 1")
    db.add(Contact(company_id=c.id, full_name="Jane", is_decision_maker=True))
    db.commit()
    r = client.patch(f"/api/companies/{c.id}", json={"lead_status": "contacted"})
    assert r.status_code == 200
    body = r.json()
    assert body["lead_status"] == "contacted"
    assert body["score"] is not None and body["score_reasons"]
    assert body["contacts"][0]["full_name"] == "Jane"

    # the same status again logs nothing
    client.patch(f"/api/companies/{c.id}", json={"lead_status": "contacted"})
    client.patch(f"/api/companies/{c.id}", json={"lead_status": "lost"})
    acts = client.get(f"/api/companies/{c.id}/activities").json()
    assert [a["text"] for a in acts] == ["Status: Contacted → Lost", "Status: New → Contacted"]
    assert all(a["kind"] == "status" for a in acts)
    assert client.get(f"/api/companies/{c.id}").json()["score"] == 0  # lost -> 0

    assert client.patch(f"/api/companies/{c.id}", json={"lead_status": "maybe"}).status_code == 422
    assert client.patch("/api/companies/9999", json={"lead_status": "won"}).status_code == 404


def test_notes_and_timeline_newest_first(client, db):
    c = add_company(db, "Acme")
    other = add_company(db, "Other")
    jane = Contact(company_id=c.id, full_name="Jane")
    stranger = Contact(company_id=other.id, full_name="Stranger")
    db.add_all([jane, stranger])
    db.commit()

    r = client.post(f"/api/companies/{c.id}/notes", json={"text": "  Called, call back Monday  "})
    assert r.status_code == 201
    note = r.json()
    assert (note["kind"], note["text"], note["contact_id"], note["company_id"]) == (
        "note", "Called, call back Monday", None, c.id)
    client.post(f"/api/companies/{c.id}/notes", json={"text": "Spoke to Jane", "contact_id": jane.id})

    acts = client.get(f"/api/companies/{c.id}/activities").json()
    assert [a["text"] for a in acts] == ["Spoke to Jane", "Called, call back Monday"]
    assert acts[0]["contact_id"] == jane.id

    bad_contact = client.post(f"/api/companies/{c.id}/notes", json={"text": "x", "contact_id": stranger.id})
    assert bad_contact.status_code == 400
    assert client.post(f"/api/companies/{c.id}/notes", json={"text": "   "}).status_code == 422
    assert client.post(f"/api/companies/{c.id}/notes", json={"text": ""}).status_code == 422
    assert client.post("/api/companies/9999/notes", json={"text": "x"}).status_code == 404
    assert client.get("/api/companies/9999/activities").status_code == 404
    assert client.get(f"/api/companies/{other.id}/activities").json() == []
    assert db.query(Activity).count() == 2


# ----- exports -----------------------------------------------------------------------


def test_export_xlsx_and_csv_match(client, sample):
    r = client.get("/api/companies/export.xlsx?min_score=50&sort=score")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="leads.xlsx"' in r.headers["content-disposition"]
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    header = list(rows[0])
    assert header[0] == "name" and "lead_status" in header and "score" in header
    assert [row[0] for row in rows[1:]] == ["Beta", "Delta"]

    r = client.get("/api/companies/export.csv?min_score=50&sort=score")
    assert r.headers["content-type"].startswith("text/csv")
    text = r.content.decode("utf-8")
    assert text.startswith("﻿")
    csv_rows = list(csv.reader(io.StringIO(text[1:])))
    assert csv_rows[0] == header  # same columns in both formats
    assert [row[0] for row in csv_rows[1:]] == ["Beta", "Delta"]


# ----- find emails -------------------------------------------------------------------


def test_find_emails_for_selected_companies(client, sample):
    alpha, beta, *_ = sample
    r = client.post("/api/companies/find-emails", json={"company_ids": [beta.id, alpha.id, 9999], "crawl": True})
    assert r.status_code == 201
    job = r.json()
    assert (job["kind"], job["source"], job["status"], job["parent_job_id"]) == ("emails", "emails", "queued", None)
    assert job["query"] == "2 selected companies"
    assert job["params"] == {"company_ids": [alpha.id, beta.id], "force": False, "crawl": True, "use_hunter": True}

    r = client.post("/api/companies/find-emails", json={"company_ids": [alpha.id], "force": True, "use_hunter": False})
    assert r.json()["query"] == "1 selected company"
    assert r.json()["params"] == {"company_ids": [alpha.id], "force": True, "crawl": False, "use_hunter": False}

    assert client.post("/api/companies/find-emails", json={"company_ids": [9999]}).status_code == 400
    assert client.post("/api/companies/find-emails", json={"company_ids": []}).status_code == 422


def test_score_filters_must_be_0_to_100(client, db, sample):
    """A huge min_intent / min_score used to crash the database query (a 500 error)."""
    from app.routers.companies import filters

    huge = 10**23
    for name in ("min_intent", "min_score"):
        assert client.get(f"/api/companies?{name}={huge}").status_code == 422
        assert client.get(f"/api/companies/export.csv?{name}={huge}").status_code == 422
        assert client.get(f"/api/companies?{name}=101").status_code == 422
        assert client.get(f"/api/companies?{name}=-1").status_code == 422
    assert client.get("/api/companies?min_intent=60").status_code == 200
    # The Google Sheets export calls filters() as a plain function: checked there too.
    for bad in (huge, "abc", 101, -1, True):
        with pytest.raises(ValueError):
            filters(min_intent=bad)
        with pytest.raises(ValueError):
            filters(min_score=bad)
    filters(min_intent=0, min_score=100)  # fine
    resp = client.post("/api/export/sheets", json={"kind": "companies", "filters": {"min_intent": huge},
                                                   "spreadsheet": "1" * 30})
    assert resp.status_code == 400 and "min_intent" in resp.text
